#!/usr/bin/env python3
"""DOCX formatting engines used by the native macOS client.

The default path uses python-docx for document persistence and the bundled
OOXML helpers for equations and table styles.  The original standard-library
OOXML formatter remains available as an explicit fallback engine.
"""

from __future__ import annotations

import json
import io
import os
import re
import sys
import tempfile
import zipfile
from copy import deepcopy
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.etree import ElementTree as ET

# The engine runs from a signed App bundle; never write import caches beside it.
sys.dont_write_bytecode = True


BUNDLED_PYTHON_PACKAGES = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "python_packages"
)
if os.path.isdir(BUNDLED_PYTHON_PACKAGES):
    sys.path.insert(0, BUNDLED_PYTHON_PACKAGES)


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"

NS = {"w": W_NS, "m": M_NS, "r": R_NS}

for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)


def qn(namespace: str, name: str) -> str:
    return "{%s}%s" % (namespace, name)


def w(name: str) -> str:
    return qn(W_NS, name)


def m(name: str) -> str:
    return qn(M_NS, name)


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def namespace_uri(name: str) -> str:
    if name.startswith("{") and "}" in name:
        return name[1:].split("}", 1)[0]
    return ""


def register_namespaces(data: bytes) -> Dict[str, str]:
    """Register source prefixes before ElementTree serializes a part again."""
    mappings: Dict[str, str] = {}
    for _, (prefix, uri) in ET.iterparse(io.BytesIO(data), events=("start-ns",)):
        prefix = prefix or ""
        mappings.setdefault(prefix, uri)
        try:
            ET.register_namespace(prefix, uri)
        except ValueError:
            # Invalid or reserved source prefixes are not needed for output.
            pass
    return mappings


def sanitize_ignorable_namespaces(root: ET.Element, mappings: Dict[str, str]) -> None:
    """Remove mc:Ignorable prefixes whose declarations were not serialized."""
    used_uris = set()
    for element in root.iter():
        used_uris.add(namespace_uri(element.tag))
        used_uris.update(namespace_uri(name) for name in element.attrib)

    ignorable = qn(MC_NS, "Ignorable")
    for element in root.iter():
        if ignorable not in element.attrib:
            continue
        prefixes = str(element.attrib[ignorable]).split()
        retained = [
            prefix
            for prefix in prefixes
            if mappings.get(prefix, "") in used_uris
        ]
        if retained:
            element.set(ignorable, " ".join(retained))
        else:
            del element.attrib[ignorable]


def serialize_xml(root: ET.Element, mappings: Dict[str, str]) -> bytes:
    sanitize_ignorable_namespaces(root, mappings)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def child(parent: ET.Element, tag: str) -> Optional[ET.Element]:
    return parent.find(tag, NS)


def ensure_child(parent: ET.Element, tag: str, index: Optional[int] = None) -> ET.Element:
    existing = child(parent, tag)
    if existing is not None:
        return existing
    element = ET.Element(tag)
    if index is None:
        parent.append(element)
    else:
        parent.insert(index, element)
    return element


def remove_child(parent: ET.Element, tag: str) -> None:
    existing = child(parent, tag)
    if existing is not None:
        parent.remove(existing)


def set_attr(element: ET.Element, name: str, value: Any) -> None:
    element.set(qn(W_NS, name), str(value))


def to_float(value: Any, fallback: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return fallback


def to_int(value: Any, fallback: int) -> int:
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return fallback


def points_to_twips(value: Any) -> int:
    return max(0, to_int(value, 0) * 20)


def cm_to_twips(value: Any) -> int:
    return max(0, to_int(to_float(value, 2.54) * 566.929134, 1440))


def half_points(value: Any) -> int:
    return max(2, to_int(to_float(value, 12) * 2, 24))


def default_style(
    font_name: str = "Times New Roman",
    font_size: float = 12,
    line_spacing: float = 1.5,
    before: float = 0,
    after: float = 0,
    first_line_indent: float = 0,
    alignment: str = "left",
    bold: bool = False,
    italic: bool = False,
    latin_font_name: Optional[str] = None,
) -> Dict[str, Any]:
    style = {
        "font_name": font_name,
        "font_size": font_size,
        "line_spacing": line_spacing,
        "before": before,
        "after": after,
        "first_line_indent": first_line_indent,
        "alignment": alignment,
        "bold": bold,
        "italic": italic,
    }
    if latin_font_name:
        style["latin_font_name"] = latin_font_name
    return style


def default_config() -> Dict[str, Any]:
    return {
        "page": {
            "paper_size": "LETTER",
            "margin_top": 2.54,
            "margin_bottom": 2.54,
            "margin_left": 3.18,
            "margin_right": 3.18,
        },
        "body": default_style(
            font_name="宋体",
            font_size=12,
            line_spacing=1.15,
            before=9,
            after=9,
            first_line_indent=24,
            alignment="justify",
            latin_font_name="Times New Roman",
        ),
        "heading1": default_style(
            font_name="黑体",
            font_size=20,
            line_spacing=1.15,
            before=18,
            after=4,
            alignment="center",
            bold=True,
        ),
        "heading2": default_style(
            font_name="黑体",
            font_size=16,
            line_spacing=1.15,
            before=36,
            after=0,
            alignment="center",
            bold=True,
        ),
        "heading3": default_style(
            font_name="宋体",
            font_size=14,
            line_spacing=1.15,
            before=6,
            after=0,
            bold=True,
        ),
        "heading4": default_style(
            font_name="宋体",
            font_size=12,
            line_spacing=1.15,
            before=4,
            after=2,
            bold=True,
        ),
        "heading5": default_style(
            font_name="宋体",
            font_size=11,
            line_spacing=1.15,
            before=4,
            after=2,
            bold=True,
        ),
        "heading6": default_style(
            font_name="宋体",
            font_size=10,
            line_spacing=1.15,
            before=2,
            after=0,
            bold=True,
        ),
        "equation": default_style(
            font_name="Latin Modern Math",
            font_size=11,
            line_spacing=1.15,
            before=9,
            after=9,
            alignment="center",
        ),
        "image": default_style(
            font_name="宋体",
            font_size=12,
            line_spacing=1.0,
            before=9,
            after=9,
            first_line_indent=0,
            alignment="center",
        ),
        "table": {
            **default_style(
                font_name="宋体",
                font_size=9,
                line_spacing=1.5,
                alignment="center",
                latin_font_name="Latin Modern Math",
            ),
            "style_id": "af2",
            "width_mode": "percentage",
            "width_percent": 100,
            "auto_width_max_columns": 3,
            "paragraph_style_id": "Compact",
            "header_bold": True,
        },
        "caption": default_style(
            font_name="宋体",
            font_size=10,
            line_spacing=1.15,
            after=4,
            alignment="center",
            bold=True,
        ),
        "note": default_style(
            font_name="宋体",
            font_size=10,
            line_spacing=1.0,
            before=0,
            after=4,
            first_line_indent=0,
            alignment="left",
            latin_font_name="Times New Roman",
        ),
        "reference": default_style(
            font_name="宋体",
            font_size=12,
            line_spacing=1.15,
            before=9,
            after=9,
            first_line_indent=0,
            alignment="left",
            latin_font_name="Times New Roman",
        ),
        "code": {
            **default_style(
                font_name="Menlo",
                font_size=9,
                line_spacing=1.0,
                before=4,
                after=4,
                first_line_indent=0,
                alignment="left",
            ),
            "preserve_inline_styles": False,
        },
        "list": default_style(
            font_name="宋体",
            font_size=12,
            line_spacing=1.15,
            before=0,
            after=0,
            first_line_indent=0,
            alignment="left",
        ),
    }


def deep_merge(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    merged = deepcopy(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = deep_merge(merged[key], value)
        else:
            merged[key] = value
    return merged


def text_content(element: ET.Element) -> str:
    chunks: List[str] = []
    for node in element.iter():
        name = local_name(node.tag)
        if name in ("t", "instrText") and node.text:
            chunks.append(node.text)
        elif name == "tab":
            chunks.append("\t")
        elif name in ("br", "cr"):
            chunks.append("\n")
    return "".join(chunks).strip()


def paragraph_style_id(paragraph: ET.Element) -> str:
    ppr = child(paragraph, "w:pPr")
    if ppr is None:
        return ""
    style = child(ppr, "w:pStyle")
    if style is None:
        return ""
    return style.get(qn(W_NS, "val"), "")


def build_style_catalog(styles_root: Optional[ET.Element]) -> Dict[str, Dict[str, Any]]:
    catalog: Dict[str, Dict[str, Any]] = {}
    if styles_root is None:
        return catalog
    for style in styles_root.findall("w:style", NS):
        style_id = style.get(qn(W_NS, "styleId"), "")
        if not style_id:
            continue
        name_node = child(style, "w:name")
        based_node = child(style, "w:basedOn")
        style_ppr = child(style, "w:pPr")
        outline_node = child(style_ppr, "w:outlineLvl") if style_ppr is not None else None
        catalog[style_id] = {
            "type": style.get(qn(W_NS, "type"), ""),
            "name": name_node.get(qn(W_NS, "val"), "") if name_node is not None else "",
            "based_on": based_node.get(qn(W_NS, "val"), "") if based_node is not None else "",
            "outline": to_int(outline_node.get(qn(W_NS, "val")), 99) if outline_node is not None else None,
        }
    return catalog


def style_info(paragraph: ET.Element, styles: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    return styles.get(paragraph_style_id(paragraph), {})


def has_math(paragraph: ET.Element) -> bool:
    return any(local_name(node.tag) in ("oMath", "oMathPara") for node in paragraph.iter())


def has_block_math(paragraph: ET.Element) -> bool:
    return any(local_name(node.tag) == "oMathPara" for node in paragraph.iter())


def ordinary_text(paragraph: ET.Element) -> str:
    chunks: List[str] = []
    for node in paragraph.iter():
        if node.tag == w("t") and node.text:
            chunks.append(node.text)
        elif node.tag in (w("tab"),):
            chunks.append("\t")
        elif node.tag in (w("br"), w("cr")):
            chunks.append("\n")
    return "".join(chunks).strip()


def has_drawing(paragraph: ET.Element) -> bool:
    return any(local_name(node.tag) in ("drawing", "pict", "object") for node in paragraph.iter())


def normalized(value: str) -> str:
    return re.sub(r"[\s_\-]", "", value).lower()


def heading_level(paragraph: ET.Element, styles: Dict[str, Dict[str, Any]]) -> Optional[int]:
    style_id = paragraph_style_id(paragraph)
    info = style_info(paragraph, styles)
    style = normalized(style_id)
    style_name = normalized(str(info.get("name", "")))
    for level in (1, 2, 3):
        if any(
            token in style or token in style_name
            for token in (f"heading{level}", f"标题{level}", f"title{level}")
        ):
            return level

    ppr = child(paragraph, "w:pPr")
    if ppr is not None:
        outline = child(ppr, "w:outlineLvl")
        if outline is not None:
            level = to_int(outline.get(qn(W_NS, "val")), 99) + 1
            if level in (1, 2, 3):
                return level
    style_outline = info.get("outline")
    if style_outline is not None and style_outline + 1 in (1, 2, 3, 4, 5, 6):
        return style_outline + 1

    text = text_content(paragraph)
    if re.match(r"^第[一二三四五六七八九十百千]+章(?:\s|、|：|:|$)", text):
        return 1
    if re.match(r"^[一二三四五六七八九十百千]+[、.．:：]\s*[^。！？!?]{1,60}$", text):
        return 2
    numeric = re.match(r"^(\d+(?:\.\d+){0,5})[、.．]?\s*[^。！？!?]{1,60}$", text)
    if numeric:
        return min(6, numeric.group(1).count(".") + 2)
    return None


CAPTION_LABEL_PATTERN = re.compile(
    r"^\s*(?:table|figure|图|表)\s*"
    r"(?:[A-Za-z]?\d+(?:[.-]\d+)*|[一二三四五六七八九十百千]+)"
    r"(?P<rest>.*)$",
    re.IGNORECASE,
)
CAPTION_NARRATIVE_PATTERN = re.compile(
    r"^(?:中|里|内|前|后|随后|揭示|显示|表明|说明|反映|体现|展示|描述|"
    r"给出|列示|列出|报告|呈现|表现|结果显示|数据显示|数据表明|可以看出|"
    r"可见|如下|见表|见图|包含|包括|分析|比较)",
    re.IGNORECASE,
)


def is_caption(text: str) -> bool:
    """Recognize standalone table/figure captions without matching prose.

    Pandoc commonly emits captions such as ``表 1 样本划分`` without a colon,
    so whitespace-only separators remain supported.  A paragraph like
    ``图 4 中，...`` is prose, however, and must not be treated as a caption.
    """
    normalized_text = re.sub(r"\s+", " ", str(text or "").strip())
    match = CAPTION_LABEL_PATTERN.match(normalized_text)
    if match is None:
        return False

    rest = match.group("rest")
    if not rest:
        return False

    explicit_separator = re.match(r"^\s*[.．:：、\-—]\s*(?P<title>.+)$", rest)
    if explicit_separator is not None:
        title = explicit_separator.group("title").strip()
    else:
        whitespace_separator = re.match(r"^\s+(?P<title>.+)$", rest)
        if whitespace_separator is None:
            # No punctuation or whitespace after the label means this is an
            # inline expression such as ``表5揭示了…``, not a caption.
            return False
        title = whitespace_separator.group("title").strip()

        # Whitespace-only captions are inherently ambiguous.  Keep the
        # compact, title-like form used by Pandoc, but reject sentence-like
        # prose with clauses or sentence punctuation.
        if len(title) > 100 or re.search(r"[，,；;。！？!?]", title):
            return False

    if not title or CAPTION_NARRATIVE_PATTERN.match(title):
        return False
    return True


def is_note(text: str) -> bool:
    return bool(
        re.match(
            r"^\s*(?:注|注释|说明|资料来源|来源|note|notes|source|sources)\s*[:：.]",
            text,
            re.IGNORECASE,
        )
    )


def is_reference_heading(text: str) -> bool:
    return bool(
        re.match(
            r"^\s*(?:参考文献|references?|bibliography)\s*$",
            text,
            re.IGNORECASE,
        )
    )


def has_reference_number(text: str) -> bool:
    return bool(re.match(r"^\s*(?:\[\d+\]|\d+[.、])\s*", text))


def is_reference(paragraph: ET.Element, text: str, styles: Dict[str, Dict[str, Any]]) -> bool:
    style = normalized(paragraph_style_id(paragraph))
    name = normalized(str(style_info(paragraph, styles).get("name", "")))
    if any(token in style or token in name for token in ("bibliography", "reference", "参考文献")):
        return True
    return has_reference_number(text) and len(text) > 30


def is_code_paragraph(paragraph: ET.Element, styles: Dict[str, Dict[str, Any]]) -> bool:
    style_id = normalized(paragraph_style_id(paragraph))
    name = normalized(str(style_info(paragraph, styles).get("name", "")))
    return "sourcecode" in style_id or "sourcecode" in name or "codeblock" in name


def is_list_paragraph(paragraph: ET.Element) -> bool:
    ppr = child(paragraph, "w:pPr")
    return child(ppr, "w:numPr") is not None if ppr is not None else False


def is_block_equation(paragraph: ET.Element, styles: Dict[str, Dict[str, Any]]) -> bool:
    if has_block_math(paragraph):
        return True
    info = style_info(paragraph, styles)
    style_name = normalized(str(info.get("name", "")))
    if has_math(paragraph) and any(token in style_name for token in ("equation", "公式")):
        return True
    return has_math(paragraph) and not ordinary_text(paragraph)


def classify_paragraph(paragraph: ET.Element, styles: Dict[str, Dict[str, Any]]) -> Tuple[str, Optional[int]]:
    text = text_content(paragraph)
    level = heading_level(paragraph, styles)
    if level is not None:
        return "heading", level
    if is_block_equation(paragraph, styles):
        return "equation", None
    if is_code_paragraph(paragraph, styles):
        return "code", None
    if is_list_paragraph(paragraph):
        return "list", None
    if is_caption(text):
        return "caption", None
    if is_note(text):
        return "note", None
    if is_reference(paragraph, text, styles):
        return "reference", None
    if has_drawing(paragraph) and not text:
        return "image", None
    return "body", None


FORMAT_SCOPES = {
    "all",
    "page",
    "body",
    "heading1",
    "heading2",
    "heading3",
    "heading4",
    "heading5",
    "heading6",
    "equation",
    "table",
    "caption",
    "note",
    "reference",
    "code",
    "list",
}


def normalize_scope(value: Any) -> str:
    scope = str(value or "all").strip().lower()
    return scope if scope in FORMAT_SCOPES else "all"


def scope_matches(scope: str, kind: str, level: Optional[int]) -> bool:
    scope = normalize_scope(scope)
    if scope == "all":
        return True
    if scope == "page":
        return False
    if kind == "heading":
        return scope == f"heading{level}" if level is not None else False
    return scope == kind


def ensure_paragraph_properties(paragraph: ET.Element) -> ET.Element:
    ppr = child(paragraph, "w:pPr")
    if ppr is not None:
        return ppr
    ppr = ET.Element(w("pPr"))
    paragraph.insert(0, ppr)
    return ppr


def set_on_off(parent: ET.Element, tag: str, value: bool) -> None:
    element = child(parent, tag)
    if element is None:
        element = ET.SubElement(parent, tag)
    if value:
        element.attrib.pop(qn(W_NS, "val"), None)
    else:
        set_attr(element, "val", "0")


def apply_paragraph_style(
    paragraph: ET.Element,
    style: Dict[str, Any],
    preserve_indent: bool = False,
) -> None:
    ppr = ensure_paragraph_properties(paragraph)

    spacing = ensure_child(ppr, w("spacing"))
    if style.get("before") is not None:
        set_attr(spacing, "before", points_to_twips(style.get("before", 0)))
    if style.get("after") is not None:
        set_attr(spacing, "after", points_to_twips(style.get("after", 0)))
    if style.get("line_spacing") is not None:
        line_spacing = max(0.5, to_float(style.get("line_spacing", 1.0), 1.0))
        set_attr(spacing, "line", max(1, int(round(line_spacing * 240))))
        set_attr(spacing, "lineRule", "auto")

    if not preserve_indent and style.get("first_line_indent") is not None:
        indent = ensure_child(ppr, w("ind"))
        set_attr(indent, "firstLine", points_to_twips(style.get("first_line_indent", 0)))

    alignment = str(style.get("alignment", "left")).lower()
    if alignment in ("preserve", "inherit", ""):
        alignment = ""
    if alignment not in ("left", "center", "right", "both", "justify", ""):
        alignment = "left"
    if alignment == "justify":
        alignment = "both"
    if alignment:
        jc = ensure_child(ppr, w("jc"))
        set_attr(jc, "val", alignment)


def ensure_run_properties(run: ET.Element) -> ET.Element:
    rpr = child(run, "w:rPr")
    if rpr is not None:
        return rpr
    rpr = ET.Element(w("rPr"))
    run.insert(0, rpr)
    return rpr


def apply_run_style(run: ET.Element, style: Dict[str, Any], target_rpr: Optional[ET.Element] = None) -> None:
    rpr = target_rpr if target_rpr is not None else ensure_run_properties(run)
    font_name = str(style.get("font_name", "Times New Roman"))
    if font_name and font_name.lower() not in ("preserve", "inherit"):
        fonts = ensure_child(rpr, w("rFonts"), 0)
        latin_font_name = str(style.get("latin_font_name", font_name))
        for attr in ("ascii", "hAnsi"):
            set_attr(fonts, attr, latin_font_name)
        for attr in ("eastAsia", "cs"):
            set_attr(fonts, attr, font_name)

    if style.get("font_size") is not None:
        size = str(half_points(style.get("font_size", 12)))
        font_size = ensure_child(rpr, w("sz"))
        set_attr(font_size, "val", size)
        size_cs = ensure_child(rpr, w("szCs"))
        set_attr(size_cs, "val", size)

    # False means "do not add emphasis" so Markdown bold/italic runs survive
    # a document-wide body formatting pass.
    if bool(style.get("bold", False)):
        set_on_off(rpr, w("b"), True)
    if bool(style.get("italic", False)):
        set_on_off(rpr, w("i"), True)


def apply_math_style(paragraph: ET.Element, style: Dict[str, Any]) -> None:
    """Apply font controls to OMML runs without converting the equation."""
    for node in paragraph.iter():
        if local_name(node.tag) != "r" or not node.tag.startswith("{" + M_NS + "}"):
            continue
        math_rpr = child(node, "m:rPr")
        if math_rpr is None:
            math_rpr = ET.Element(m("rPr"))
            node.insert(0, math_rpr)
        ctrl = child(math_rpr, "m:ctrlPr")
        if ctrl is None:
            ctrl = ET.SubElement(math_rpr, m("ctrlPr"))
        word_rpr = child(ctrl, "w:rPr")
        if word_rpr is None:
            word_rpr = ET.SubElement(ctrl, w("rPr"))
        apply_run_style(node, style, target_rpr=word_rpr)


def run_style_id(run: ET.Element) -> str:
    rpr = child(run, "w:rPr")
    style = child(rpr, "w:rStyle") if rpr is not None else None
    return style.get(qn(W_NS, "val"), "") if style is not None else ""


def apply_runs(paragraph: ET.Element, style: Dict[str, Any]) -> None:
    preserve_inline = bool(style.get("preserve_inline_styles", True))
    for run in paragraph.findall(".//w:r", NS):
        token_style = normalized(run_style_id(run))
        if preserve_inline and (
            token_style in ("verbatimchar", "code")
            or token_style.endswith("tok")
            or token_style.endswith("char")
        ):
            continue
        apply_run_style(run, style)


def prepend_reference_number(paragraph: ET.Element, number: int) -> None:
    """Add a stable visible reference number without changing hyperlink runs."""
    if has_reference_number(text_content(paragraph)):
        return

    run = ET.Element(w("r"))
    text_node = ET.SubElement(run, w("t"))
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = f"[{number}] "

    ppr = child(paragraph, "w:pPr")
    insert_at = list(paragraph).index(ppr) + 1 if ppr is not None else 0
    paragraph.insert(insert_at, run)


def set_paragraph_style_id(paragraph: ET.Element, style_id: str, styles: Dict[str, Dict[str, Any]]) -> None:
    if not style_id or style_id not in styles:
        return
    ppr = ensure_paragraph_properties(paragraph)
    pstyle = ensure_child(ppr, w("pStyle"), 0)
    set_attr(pstyle, "val", style_id)


def ensure_table_style(styles_root: Optional[ET.Element], requested: str) -> str:
    if styles_root is None:
        return requested
    existing = None
    for style in styles_root.findall("w:style", NS):
        if style.get(qn(W_NS, "styleId")) == requested:
            existing = style
            break
    if existing is not None and existing.get(qn(W_NS, "type")) == "table":
        return requested

    for style in styles_root.findall("w:style", NS):
        name_node = child(style, "w:name")
        if (
            style.get(qn(W_NS, "type")) == "table"
            and name_node is not None
            and name_node.get(qn(W_NS, "val")) == "三线表"
        ):
            return style.get(qn(W_NS, "styleId"), requested)

    # Pandoc commonly uses af2 for a paragraph style named "公式".  Reusing
    # that ID as a table style would make the result dependent on the input
    # template, so fall back to its existing three-line table base when needed.
    for style in styles_root.findall("w:style", NS):
        if style.get(qn(W_NS, "styleId")) == "temp" and style.get(qn(W_NS, "type")) == "table":
            return "temp"

    base = ET.Element(w("style"), {
        qn(W_NS, "type"): "table",
        qn(W_NS, "styleId"): "PaperFormatterTable",
    })
    name = ET.SubElement(base, w("name"))
    set_attr(name, "val", "Paper Formatter Table")
    styles_root.append(base)
    return "PaperFormatterTable"


def table_math_style(row_style: Dict[str, Any]) -> Dict[str, Any]:
    """Use the table row size for math while keeping the math typeface explicit."""
    style = dict(row_style)
    style["font_name"] = "Latin Modern Math"
    style["latin_font_name"] = "Latin Modern Math"
    if row_style.get("font_size") is not None:
        style["font_size"] = row_style["font_size"]
    return style


def apply_table_properties(
    table: ET.Element,
    config: Dict[str, Any],
    styles: Dict[str, Dict[str, Any]],
    styles_root: Optional[ET.Element],
) -> None:
    tblpr = child(table, "w:tblPr")
    if tblpr is None:
        tblpr = ET.Element(w("tblPr"))
        table.insert(0, tblpr)

    requested_style = str(config.get("style_id", "af2"))
    table_style = ensure_table_style(styles_root, requested_style)
    tbl_style = ensure_child(tblpr, w("tblStyle"), 0)
    set_attr(tbl_style, "val", table_style)

    rows = table.findall("w:tr", NS)
    column_count = max((len(row.findall("w:tc", NS)) for row in rows), default=0)
    width_mode = str(config.get("width_mode", "auto_by_columns")).lower()
    if width_mode not in ("preserve", "inherit"):
        tbl_width = ensure_child(tblpr, w("tblW"))
        use_auto = width_mode == "auto" or (
            width_mode in ("auto_by_columns",)
            and column_count <= to_int(config.get("auto_width_max_columns", 3), 3)
        )
        if use_auto:
            set_attr(tbl_width, "w", 0)
            set_attr(tbl_width, "type", "auto")
        else:
            percent = max(1, min(100, to_float(config.get("width_percent", 50), 50)))
            set_attr(tbl_width, "w", int(round(percent * 50)))
            set_attr(tbl_width, "type", "pct")

    look = ensure_child(tblpr, w("tblLook"))
    for name, value in {
        "val": "04A0",
        "firstRow": "1",
        "lastRow": "0",
        "firstColumn": "1",
        "lastColumn": "0",
        "noHBand": "0",
        "noVBand": "1",
    }.items():
        set_attr(look, name, value)


def apply_table(
    table: ET.Element,
    config: Dict[str, Any],
    styles: Dict[str, Dict[str, Any]],
    styles_root: Optional[ET.Element],
) -> None:
    apply_table_properties(table, config, styles, styles_root)
    paragraph_style = str(config.get("paragraph_style_id", "Compact"))
    header_style = dict(config)
    if bool(config.get("header_bold", True)):
        header_style["bold"] = True

    rows = table.findall("w:tr", NS)
    for row_index, row in enumerate(rows):
        row_config = header_style if row_index == 0 else config
        for paragraph in row.findall(".//w:p", NS):
            set_paragraph_style_id(paragraph, paragraph_style, styles)
            apply_paragraph_style(paragraph, row_config, preserve_indent=True)
            apply_runs(paragraph, row_config)
            if has_math(paragraph):
                apply_math_style(paragraph, table_math_style(row_config))


def apply_page_settings(root: ET.Element, page: Dict[str, Any]) -> None:
    body = root.find(".//w:body", NS)
    if body is None:
        return
    sect = body.find("w:sectPr", NS)
    if sect is None:
        sect = ET.SubElement(body, w("sectPr"))
    page_size = ensure_child(sect, w("pgSz"))
    paper = str(page.get("paper_size", "A4")).upper()
    if paper == "LETTER":
        set_attr(page_size, "w", 12240)
        set_attr(page_size, "h", 15840)
    else:
        set_attr(page_size, "w", 11906)
        set_attr(page_size, "h", 16838)
    margins = ensure_child(sect, w("pgMar"))
    set_attr(margins, "top", cm_to_twips(page.get("margin_top", 2.54)))
    set_attr(margins, "bottom", cm_to_twips(page.get("margin_bottom", 2.54)))
    set_attr(margins, "left", cm_to_twips(page.get("margin_left", 3.18)))
    set_attr(margins, "right", cm_to_twips(page.get("margin_right", 3.18)))


def document_body(root: ET.Element) -> Optional[ET.Element]:
    return root.find(".//w:body", NS)


def walk_blocks(root: ET.Element) -> Iterable[Tuple[str, ET.Element]]:
    body = document_body(root)
    if body is None:
        return
    for block in list(body):
        name = local_name(block.tag)
        if name == "p":
            yield "paragraph", block
        elif name == "tbl":
            yield "table", block


def inspect_tree(root: ET.Element, styles: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    counts = {
        "paragraphs": 0,
        "tables": 0,
        "equations": 0,
        "images": 0,
        "headings": 0,
        "captions": 0,
        "notes": 0,
        "references": 0,
        "code_blocks": 0,
        "lists": 0,
    }
    samples: List[Dict[str, str]] = []
    in_reference_section = False
    for block_type, block in walk_blocks(root):
        if block_type == "table":
            counts["tables"] += 1
            if len(samples) < 12:
                samples.append({"type": "table", "label": "表格", "text": "表格对象"})
            continue

        counts["paragraphs"] += 1
        text = text_content(block).strip()
        kind, level = classify_paragraph(block, styles)
        reference_heading = is_reference_heading(text)
        if reference_heading and level is None:
            kind, level = "heading", 1
        if in_reference_section and text and not reference_heading and level is not None:
            in_reference_section = False
        if in_reference_section and text and not reference_heading:
            kind, level = "reference", None
        if kind == "equation":
            counts["equations"] += 1
        elif kind == "image":
            counts["images"] += 1
        elif kind == "heading":
            counts["headings"] += 1
        elif kind == "code":
            counts["code_blocks"] += 1
        elif kind == "list":
            counts["lists"] += 1
        elif kind == "caption":
            counts["captions"] += 1
        elif kind == "note":
            counts["notes"] += 1
        elif kind == "reference":
            counts["references"] += 1
        if len(samples) < 12 and kind != "body":
            label = kind if level is None else f"heading{level}"
            samples.append({"type": kind, "label": label, "text": text[:100]})
        if reference_heading:
            in_reference_section = True
    counts["blocks"] = counts["paragraphs"] + counts["tables"]
    return {"counts": counts, "samples": samples}


def format_tree(
    root: ET.Element,
    config: Dict[str, Any],
    styles: Dict[str, Dict[str, Any]],
    styles_root: Optional[ET.Element],
    apply_page: bool = True,
    warn_if_no_equation: bool = True,
    scope: str = "all",
) -> Dict[str, Any]:
    scope = normalize_scope(scope)
    counts = {
        "paragraphs": 0,
        "tables": 0,
        "equations": 0,
        "images": 0,
        "headings": 0,
        "captions": 0,
        "notes": 0,
        "references": 0,
        "code_blocks": 0,
        "lists": 0,
        "blocks": 0,
        "formatted": 0,
    }
    warnings: List[str] = []
    in_reference_section = False
    reference_number = 0

    if apply_page and scope in ("all", "page"):
        apply_page_settings(root, config.get("page", {}))
    for block_type, block in walk_blocks(root):
        if block_type == "table":
            counts["tables"] += 1
            if scope in ("all", "table"):
                apply_table(block, config.get("table", {}), styles, styles_root)
                counts["formatted"] += 1
            continue

        counts["paragraphs"] += 1
        text = text_content(block).strip()
        kind, level = classify_paragraph(block, styles)
        reference_heading = is_reference_heading(text)
        if in_reference_section and text and not reference_heading and level is not None:
            in_reference_section = False
        if in_reference_section and text and not reference_heading:
            kind, level = "reference", None
        selected = scope_matches(scope, kind, level)
        if kind == "equation":
            counts["equations"] += 1
        elif kind == "image":
            counts["images"] += 1
        elif kind == "heading":
            counts["headings"] += 1
        elif kind == "code":
            counts["code_blocks"] += 1
        elif kind == "list":
            counts["lists"] += 1
        elif kind == "caption":
            counts["captions"] += 1
        elif kind == "note":
            counts["notes"] += 1
        elif kind == "reference":
            counts["references"] += 1

        if selected:
            if kind == "equation":
                style = config.get("equation", {})
                apply_paragraph_style(block, style)
                apply_runs(block, style)
                apply_math_style(block, style)
            elif kind == "heading":
                style = config.get(f"heading{level}", config.get("heading3", config.get("heading1", {})))
                apply_paragraph_style(block, style)
                apply_runs(block, style)
            elif kind == "code":
                style = config.get("code", {})
                apply_paragraph_style(block, style)
                apply_runs(block, style)
            elif kind == "list":
                list_style = config.get("list", config.get("body", {}))
                apply_paragraph_style(block, list_style, preserve_indent=True)
                apply_runs(block, list_style)
            elif kind == "caption":
                caption_style = config.get("caption", {})
                apply_paragraph_style(block, caption_style)
                apply_runs(block, caption_style)
            elif kind == "note":
                note_style = config.get("note", config.get("caption", {}))
                apply_paragraph_style(block, note_style)
                apply_runs(block, note_style)
            elif kind == "reference":
                reference_number += 1
                reference_style = config.get("reference", {})
                apply_paragraph_style(block, reference_style)
                prepend_reference_number(block, reference_number)
                apply_runs(block, reference_style)
            elif kind == "image":
                apply_paragraph_style(block, config.get("image", {}))
            else:
                body_style = config.get("body", {})
                apply_paragraph_style(block, body_style)
                apply_runs(block, body_style)
            counts["formatted"] += 1

        # Inline OMML is part of the surrounding paragraph and is not a
        # separate block kind.  Allow an equation-only pass to update just
        # those math runs without changing the paragraph's body style.
        if kind != "equation" and has_math(block) and scope in ("all", "equation"):
            apply_math_style(block, config.get("equation", {}))
            if scope == "equation" and not selected:
                counts["formatted"] += 1
        if reference_heading:
            in_reference_section = True
            reference_number = 0

    if warn_if_no_equation and scope in ("all", "equation") and counts["equations"] == 0:
        warnings.append("没有检测到 Office Math 公式；图片形式的公式不会被自动识别。")
    counts["blocks"] = counts["paragraphs"] + counts["tables"]
    return {"counts": counts, "warnings": warnings, "scope": scope}


def auxiliary_story_names(entries: Dict[str, bytes]) -> List[str]:
    names = []
    for name in entries:
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        if re.match(r"^word/(header\d+|footer\d+|footnotes|endnotes|comments)\.xml$", name):
            names.append(name)
    return names


def format_auxiliary_story(
    root: ET.Element,
    config: Dict[str, Any],
    styles: Dict[str, Dict[str, Any]],
    styles_root: Optional[ET.Element],
    scope: str = "all",
) -> None:
    scope = normalize_scope(scope)
    story_style = config.get("note", config.get("reference", {}))
    for paragraph in root.findall(".//w:p", NS):
        kind, level = classify_paragraph(paragraph, styles)
        if scope_matches(scope, kind, level):
            if kind == "equation":
                style = config.get("equation", {})
                apply_paragraph_style(paragraph, style)
                apply_runs(paragraph, style)
                apply_math_style(paragraph, style)
            elif kind == "heading":
                style = config.get(f"heading{level}", config.get("heading3", {}))
                apply_paragraph_style(paragraph, style)
                apply_runs(paragraph, style)
            elif kind == "image":
                apply_paragraph_style(paragraph, config.get("image", {}), preserve_indent=True)
            elif kind == "note":
                note_style = config.get("note", story_style)
                apply_paragraph_style(paragraph, note_style, preserve_indent=True)
                apply_runs(paragraph, note_style)
            else:
                apply_paragraph_style(paragraph, story_style, preserve_indent=True)
                apply_runs(paragraph, story_style)
        if kind != "equation" and has_math(paragraph) and scope in ("all", "equation"):
            apply_math_style(paragraph, config.get("equation", {}))
    for table in root.findall(".//w:tbl", NS):
        if scope in ("all", "table"):
            apply_table(table, config.get("table", {}), styles, styles_root)


def read_docx(
    path: str,
) -> Tuple[Dict[str, bytes], ET.Element, Optional[ET.Element], Dict[str, str]]:
    if not os.path.isfile(path):
        raise FileNotFoundError(f"输入文件不存在：{path}")
    with zipfile.ZipFile(path, "r") as archive:
        entries = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    if "word/document.xml" not in entries:
        raise ValueError("这不是有效的 DOCX 文件：缺少 word/document.xml。")

    namespace_mappings: Dict[str, str] = {}
    for name, data in entries.items():
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        try:
            for prefix, uri in register_namespaces(data).items():
                namespace_mappings.setdefault(prefix, uri)
        except ET.ParseError:
            # The main document and styles parts below report their own errors.
            continue

    try:
        root = ET.fromstring(entries["word/document.xml"])
    except ET.ParseError as exc:
        raise ValueError(f"无法解析 Word 文档 XML：{exc}") from exc
    styles_root = None
    if "word/styles.xml" in entries:
        try:
            styles_root = ET.fromstring(entries["word/styles.xml"])
        except ET.ParseError as exc:
            raise ValueError(f"无法解析 Word 样式 XML：{exc}") from exc
    return entries, root, styles_root, namespace_mappings


def write_docx(
    path: str,
    entries: Dict[str, bytes],
    root: ET.Element,
    styles_root: Optional[ET.Element] = None,
    namespace_mappings: Optional[Dict[str, str]] = None,
) -> None:
    output_dir = os.path.dirname(os.path.abspath(path))
    os.makedirs(output_dir, exist_ok=True)
    mappings = namespace_mappings or {}
    entries["word/document.xml"] = serialize_xml(root, mappings)
    if styles_root is not None:
        entries["word/styles.xml"] = serialize_xml(styles_root, mappings)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename, data in entries.items():
            archive.writestr(filename, data)


def atomic_write_docx(output_path: str, writer: Any) -> None:
    """Write a DOCX beside its destination, then atomically replace it.

    This is especially important when the selected output is the input file:
    a failed formatter run must leave the original document intact.
    """
    target_path = os.path.realpath(os.path.abspath(output_path))
    output_dir = os.path.dirname(target_path)
    os.makedirs(output_dir, exist_ok=True)
    file_descriptor, staged_path = tempfile.mkstemp(
        prefix=f".{os.path.basename(target_path)}.",
        suffix=".docx",
        dir=output_dir,
    )
    os.close(file_descriptor)
    try:
        writer(staged_path)
        os.replace(staged_path, target_path)
    finally:
        if os.path.exists(staged_path):
            os.unlink(staged_path)


def inspect_docx(path: str) -> Dict[str, Any]:
    _, root, styles_root, _ = read_docx(path)
    return inspect_tree(root, build_style_catalog(styles_root))


def format_docx_native(
    input_path: str,
    output_path: str,
    raw_config: Dict[str, Any],
    scope: str = "all",
) -> Dict[str, Any]:
    scope = normalize_scope(scope)
    config = deep_merge(default_config(), raw_config or {})
    entries, root, styles_root, namespace_mappings = read_docx(input_path)
    styles = build_style_catalog(styles_root)
    result = format_tree(root, config, styles, styles_root, scope=scope)

    # Pandoc documents may contain headers, footnotes, endnotes and comments
    # outside the main document story.  Format those parts without applying
    # page-section settings to them.
    for name in auxiliary_story_names(entries):
        try:
            story_root = ET.fromstring(entries[name])
            format_auxiliary_story(story_root, config, styles, styles_root, scope=scope)
            entries[name] = serialize_xml(story_root, namespace_mappings)
        except ET.ParseError:
            result.setdefault("warnings", []).append(f"跳过无法解析的辅助文档流：{name}")
    atomic_write_docx(
        output_path,
        lambda staged_path: write_docx(
            staged_path, entries, root, styles_root, namespace_mappings
        ),
    )
    result["output_path"] = os.path.abspath(output_path)
    return result


_PYTHON_DOCX_API: Optional[Tuple[Any, ...]] = None


def python_docx_api() -> Tuple[Any, ...]:
    global _PYTHON_DOCX_API
    if _PYTHON_DOCX_API is not None:
        return _PYTHON_DOCX_API
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import qn as docx_qn
        from docx.shared import Cm, Pt
    except ImportError as exc:
        raise RuntimeError(
            "默认 python-docx 排版引擎不可用。请重新运行构建脚本，"
            "或安装 python-docx 及其依赖。"
        ) from exc
    _PYTHON_DOCX_API = (
        Document,
        WD_TABLE_ALIGNMENT,
        WD_ALIGN_PARAGRAPH,
        OxmlElement,
        parse_xml,
        docx_qn,
        Cm,
        Pt,
    )
    return _PYTHON_DOCX_API


def python_qn(name: str) -> str:
    return python_docx_api()[5](name)


def python_element(tag: str) -> Any:
    if tag.startswith("{") and "}" in tag:
        namespace, local = tag[1:].split("}", 1)
        prefixes = {
            W_NS: "w",
            M_NS: "m",
            R_NS: "r",
            MC_NS: "mc",
        }
        tag = f"{prefixes.get(namespace, 'w')}:{local}"
    return python_docx_api()[3](tag)


def python_child(parent: Any, tag: str) -> Optional[Any]:
    return parent.find(tag)


def python_ensure_child(parent: Any, tag: str, index: Optional[int] = None) -> Any:
    existing = python_child(parent, tag)
    if existing is not None:
        return existing
    created = python_element(tag)
    if index is None:
        parent.append(created)
    else:
        parent.insert(index, created)
    return created


def python_set_font_rpr(rpr: Any, style: Dict[str, Any]) -> None:
    font_name = str(style.get("font_name", "")).strip()
    if not font_name or font_name.lower() in ("preserve", "inherit"):
        return

    fonts = python_ensure_child(rpr, python_qn("w:rFonts"), 0)
    latin_font_name = str(style.get("latin_font_name", font_name)).strip()
    for attribute in ("ascii", "hAnsi"):
        fonts.set(python_qn(f"w:{attribute}"), latin_font_name or font_name)
    for attribute in ("eastAsia", "cs"):
        fonts.set(python_qn(f"w:{attribute}"), font_name)

    lang = python_ensure_child(rpr, python_qn("w:lang"))
    lang.set(python_qn("w:eastAsia"), "zh-CN")

    if style.get("font_size") is not None:
        half_points_value = half_points(style.get("font_size", 12))
        size = python_ensure_child(rpr, python_qn("w:sz"))
        size.set(python_qn("w:val"), str(half_points_value))
        size_cs = python_ensure_child(rpr, python_qn("w:szCs"))
        size_cs.set(python_qn("w:val"), str(half_points_value))


def python_apply_run_style(run: Any, style: Dict[str, Any]) -> None:
    _, _, _, _, _, _, _, Pt = python_docx_api()
    font_name = str(style.get("font_name", "")).strip()
    if font_name and font_name.lower() not in ("preserve", "inherit"):
        run.font.name = font_name
    if style.get("font_size") is not None:
        run.font.size = Pt(to_float(style.get("font_size"), 12))
    if bool(style.get("bold", False)):
        run.font.bold = True
    if bool(style.get("italic", False)):
        run.font.italic = True
    python_set_font_rpr(run._r.get_or_add_rPr(), style)


def python_apply_xml_run_style(run: Any, style: Dict[str, Any]) -> None:
    """Style a raw w:r, including runs nested inside hyperlinks."""
    rpr = python_ensure_child(run, python_qn("w:rPr"), 0)
    python_set_font_rpr(rpr, style)
    if bool(style.get("bold", False)):
        bold = python_ensure_child(rpr, python_qn("w:b"))
        bold.attrib.pop(python_qn("w:val"), None)
    if bool(style.get("italic", False)):
        italic = python_ensure_child(rpr, python_qn("w:i"))
        italic.attrib.pop(python_qn("w:val"), None)


def python_apply_paragraph_style(
    paragraph: Any,
    style: Dict[str, Any],
    preserve_indent: bool = False,
) -> None:
    _, _, WD_ALIGN_PARAGRAPH, _, _, _, _, Pt = python_docx_api()
    paragraph_format = paragraph.paragraph_format
    if style.get("before") is not None:
        paragraph_format.space_before = Pt(to_float(style.get("before"), 0))
    if style.get("after") is not None:
        paragraph_format.space_after = Pt(to_float(style.get("after"), 0))
    if style.get("line_spacing") is not None:
        paragraph_format.line_spacing = max(0.5, to_float(style.get("line_spacing"), 1.0))
    if not preserve_indent and style.get("first_line_indent") is not None:
        paragraph_format.first_line_indent = Pt(to_float(style.get("first_line_indent"), 0))

    alignment = str(style.get("alignment", "left")).lower()
    alignment_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    if alignment in alignment_map:
        paragraph.alignment = alignment_map[alignment]

    ppr = paragraph._p.get_or_add_pPr()
    paragraph_rpr = python_ensure_child(ppr, python_qn("w:rPr"))
    python_set_font_rpr(paragraph_rpr, style)


def python_apply_runs(paragraph: Any, style: Dict[str, Any]) -> None:
    # paragraph.runs omits w:r elements nested inside w:hyperlink.
    for run in paragraph._p.iter(python_qn("w:r")):
        python_apply_xml_run_style(run, style)


def python_prepend_reference_number(paragraph: Any, number: int) -> None:
    """Add a visible number before the first run while keeping hyperlinks intact."""
    if has_reference_number(paragraph.text):
        return

    run = python_element("w:r")
    text_node = python_element("w:t")
    text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = f"[{number}] "
    run.append(text_node)

    paragraph_element = paragraph._p
    ppr = paragraph_element.pPr
    insert_at = paragraph_element.index(ppr) + 1 if ppr is not None else 0
    paragraph_element.insert(insert_at, run)


def python_apply_math_style(paragraph: Any, style: Dict[str, Any]) -> None:
    """Keep OMML run properties Word can parse reliably.

    Word's current parser rejects the verbose ``m:ctrlPr/w:rPr`` block that
    python-docx used to receive here, while WPS silently removes it when it
    resaves the document.  The math font is configured at document level and
    on legacy direct OMML run properties, so existing mathematical styling
    such as ``m:sty`` is retained without keeping a Cambria Math override.
    """
    font_name = str(style.get("font_name", "Latin Modern Math")).strip()
    if not font_name or font_name.lower() in ("preserve", "inherit"):
        font_name = "Latin Modern Math"

    for math_root in (
        node
        for node in paragraph._p.iter()
        if node.tag in (python_qn("m:oMath"), python_qn("m:oMathPara"))
    ):
        # ctrlPr can occur below m:rPr, m:fPr, m:dPr, m:radPr and other OMML
        # structure nodes. Any one of them can override the document font.
        for parent in list(math_root.iter()):
            for child in list(parent):
                if child.tag == python_qn("m:ctrlPr"):
                    parent.remove(child)
                elif child.tag == python_qn("w:rPr"):
                    python_set_math_font_rpr(child, font_name, style.get("font_size"))

        # Some Pandoc/WPS files store a direct w:rPr beside m:rPr. Add one to
        # runs that lack it so Word has the same explicit math font everywhere.
        for math_run in math_root.iter(python_qn("m:r")):
            direct_rpr = python_child(math_run, python_qn("w:rPr"))
            if direct_rpr is None:
                direct_rpr = python_element("w:rPr")
                math_rpr = python_child(math_run, python_qn("m:rPr"))
                insert_index = math_run.index(math_rpr) + 1 if math_rpr is not None else 0
                math_run.insert(insert_index, direct_rpr)
            python_set_math_font_rpr(direct_rpr, font_name, style.get("font_size"))


def python_set_math_font_rpr(
    rpr: Any,
    font_name: str,
    font_size: Any = None,
) -> None:
    fonts = python_ensure_child(rpr, python_qn("w:rFonts"), 0)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(python_qn(f"w:{attribute}"), font_name)
    if font_size is not None:
        size_value = str(half_points(font_size))
        size = python_ensure_child(rpr, python_qn("w:sz"))
        size.set(python_qn("w:val"), size_value)
        size_cs = python_ensure_child(rpr, python_qn("w:szCs"))
        size_cs.set(python_qn("w:val"), size_value)


def python_set_document_math_font(doc: Any, style: Dict[str, Any]) -> None:
    """Set the standard document-level OMML font declaration.

    ``m:mathPr`` belongs before ``w:themeFontLang`` in ``settings.xml``.  The
    insertion order matters to Word even though the XML is otherwise valid.
    """
    settings = doc.settings.element
    math_properties = python_child(settings, python_qn("m:mathPr"))
    if math_properties is None:
        math_properties = python_element("m:mathPr")
        theme_font_language = python_child(settings, python_qn("w:themeFontLang"))
        if theme_font_language is None:
            settings.append(math_properties)
        else:
            settings.insert(settings.index(theme_font_language), math_properties)

    math_font = python_child(math_properties, python_qn("m:mathFont"))
    if math_font is None:
        math_font = python_element("m:mathFont")
        math_properties.insert(0, math_font)
    font_name = str(style.get("font_name", "Latin Modern Math")).strip()
    math_font.set(python_qn("m:val"), font_name or "Latin Modern Math")


def python_style_id(paragraph: Any) -> str:
    try:
        return str(paragraph.style.style_id or "")
    except (AttributeError, KeyError):
        return ""


def python_style_name(paragraph: Any) -> str:
    try:
        return str(paragraph.style.name or "")
    except (AttributeError, KeyError):
        return ""


def python_heading_level(paragraph: Any) -> Optional[int]:
    style_id = normalized(python_style_id(paragraph))
    style_name = normalized(python_style_name(paragraph))
    for level in range(1, 7):
        if any(
            token in style_id or token in style_name
            for token in (f"heading{level}", f"标题{level}", f"title{level}")
        ):
            return level

    ppr = python_child(paragraph._p, python_qn("w:pPr"))
    outline = python_child(ppr, python_qn("w:outlineLvl")) if ppr is not None else None
    if outline is not None:
        level = to_int(outline.get(python_qn("w:val")), 99) + 1
        if 1 <= level <= 6:
            return level

    text = paragraph.text.strip()
    if re.match(r"^第[一二三四五六七八九十百千]+章(?:\s|、|：|:|$)", text):
        return 1
    if re.match(r"^[一二三四五六七八九十百千]+[、.．:：]\s*[^。！？!?]{1,60}$", text):
        return 2
    numeric = re.match(r"^(\d+(?:\.\d+){0,5})[、.．]?\s*[^。！？!?]{1,60}$", text)
    if numeric:
        return min(6, numeric.group(1).count(".") + 2)
    return None


def python_is_equation(paragraph: Any) -> bool:
    return paragraph._p.find(".//" + python_qn("m:oMathPara")) is not None


def python_has_math(paragraph: Any) -> bool:
    return paragraph._p.find(".//" + python_qn("m:oMath")) is not None


def python_is_list(paragraph: Any) -> bool:
    ppr = python_child(paragraph._p, python_qn("w:pPr"))
    return ppr is not None and python_child(ppr, python_qn("w:numPr")) is not None


def python_is_code(paragraph: Any) -> bool:
    style = normalized(python_style_id(paragraph) + python_style_name(paragraph))
    return "sourcecode" in style or "codeblock" in style


def python_classify_paragraph(paragraph: Any) -> Tuple[str, Optional[int]]:
    level = python_heading_level(paragraph)
    if level is not None:
        return "heading", level
    if python_is_equation(paragraph):
        return "equation", None
    if python_is_code(paragraph):
        return "code", None
    if python_is_list(paragraph):
        return "list", None
    if is_caption(paragraph.text):
        return "caption", None
    if is_note(paragraph.text):
        return "note", None
    text = paragraph.text.strip()
    if re.match(r"^\s*(?:\[[0-9]+\]|[0-9]+[.、])\s*\S+", text) and len(text) > 30:
        return "reference", None
    if any(local_name(node.tag) in ("drawing", "pict", "object") for node in paragraph._p.iter()):
        if not text:
            return "image", None
    return "body", None


def python_style_template(doc: Any, requested_style_id: str) -> str:
    styles = doc.styles.element
    style_nodes = styles.findall(python_qn("w:style"))
    for style in style_nodes:
        if (
            style.get(python_qn("w:styleId")) == requested_style_id
            and style.get(python_qn("w:type")) == "table"
        ):
            return requested_style_id

    for style in style_nodes:
        name = style.find(python_qn("w:name"))
        if (
            style.get(python_qn("w:type")) == "table"
            and name is not None
            and name.get(python_qn("w:val")) == "三线表"
        ):
            return style.get(python_qn("w:styleId"), requested_style_id)

    style_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "three_line_table_style.xml")
    try:
        with open(style_path, "rb") as handle:
            template = python_docx_api()[4](handle.read())
    except OSError as exc:
        raise RuntimeError(f"找不到内置三线表样式模板：{style_path}") from exc

    style_id = "PaperFormatterThreeLine"
    for style in style_nodes:
        if style.get(python_qn("w:styleId")) == style_id:
            return style_id
    template.set(python_qn("w:styleId"), style_id)
    styles.append(template)
    return style_id


def python_set_table_width(table: Any, config: Dict[str, Any]) -> None:
    tblpr = table._tbl.tblPr
    width_mode = str(config.get("width_mode", "percentage")).lower()
    if width_mode in ("preserve", "inherit"):
        return
    tbl_width = python_ensure_child(tblpr, python_qn("w:tblW"), 0)
    if width_mode == "auto":
        tbl_width.set(python_qn("w:type"), "auto")
        tbl_width.set(python_qn("w:w"), "0")
        return
    percent = max(1, min(100, to_float(config.get("width_percent", 100), 100)))
    tbl_width.set(python_qn("w:type"), "pct")
    tbl_width.set(python_qn("w:w"), str(int(round(percent * 50))))


def python_remove_table_borders(table: Any) -> None:
    tblpr = table._tbl.tblPr
    borders = python_ensure_child(tblpr, python_qn("w:tblBorders"))
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = python_ensure_child(borders, python_qn(f"w:{border_name}"))
        border.set(python_qn("w:val"), "nil")
    for row in table.rows:
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            cell_borders = python_ensure_child(tcpr, python_qn("w:tcBorders"))
            for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = python_ensure_child(cell_borders, python_qn(f"w:{border_name}"))
                border.set(python_qn("w:val"), "nil")


def python_format_table(table: Any, config: Dict[str, Any], style_id: str) -> None:
    _, WD_TABLE_ALIGNMENT, _, _, _, _, _, _ = python_docx_api()
    tblpr = table._tbl.tblPr
    tbl_style = python_ensure_child(tblpr, python_qn("w:tblStyle"), 0)
    tbl_style.set(python_qn("w:val"), style_id)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    python_set_table_width(table, config)

    rows = list(table.rows)
    if rows:
        trpr = rows[0]._tr.get_or_add_trPr()
        header = python_ensure_child(trpr, python_qn("w:tblHeader"))
        header.set(python_qn("w:val"), "on")

    for row_index, row in enumerate(rows):
        row_config = dict(config)
        if row_index == 0 and bool(config.get("header_bold", True)):
            row_config["bold"] = True
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            vertical = python_ensure_child(tcpr, python_qn("w:vAlign"))
            vertical.set(python_qn("w:val"), "center")
            for paragraph in cell.paragraphs:
                paragraph_style_id = str(config.get("paragraph_style_id", ""))
                if paragraph_style_id:
                    ppr = paragraph._p.get_or_add_pPr()
                    pstyle = python_ensure_child(ppr, python_qn("w:pStyle"), 0)
                    pstyle.set(python_qn("w:val"), paragraph_style_id)
                python_apply_paragraph_style(paragraph, row_config, preserve_indent=True)
                python_apply_runs(paragraph, row_config)
                if python_has_math(paragraph):
                    python_apply_math_style(paragraph, table_math_style(row_config))


def python_extract_equation_number(math_para: Any) -> Optional[str]:
    math_root = math_para.find(python_qn("m:oMath"))
    if math_root is None:
        return None
    math_runs = list(math_root.iter(python_qn("m:r")))
    for run_index in range(len(math_runs) - 1, -1, -1):
        math_run = math_runs[run_index]
        text_nodes = list(math_run.iter(python_qn("m:t")))
        if not text_nodes:
            continue
        text_node = text_nodes[-1]
        text = text_node.text or ""
        match = re.search(r"[\u2001\s]*(#?\(\d+\))\s*$", text)
        if match is None:
            continue
        number = match.group(1).lstrip("#")
        prefix = text[: match.start()]
        if prefix.strip():
            text_node.text = prefix
        else:
            parent = math_run.getparent()
            if parent is not None:
                parent.remove(math_run)

        # Pandoc places two em spaces before the inline equation number.  The
        # number is moved to a separate table cell, so those spaces must not
        # remain part of the centered mathematical expression.
        separator_index = run_index - 1
        while separator_index >= 0:
            separator_run = math_runs[separator_index]
            separator_text = "".join(
                node.text or ""
                for node in separator_run.iter(python_qn("m:t"))
            )
            if not separator_text or separator_text.strip():
                break
            parent = separator_run.getparent()
            if parent is not None:
                parent.remove(separator_run)
            separator_index -= 1
        return number
    return None


def python_equation_table_width_twips(doc: Any) -> int:
    """Return the usable width of the first document section in twips."""
    try:
        section = doc.sections[0]
        page_width = int(section.page_width or 0)
        left_margin = int(section.left_margin or 0)
        right_margin = int(section.right_margin or 0)
        width_emu = page_width - left_margin - right_margin
        if width_emu > 0:
            return max(1, int(round(width_emu / 635)))
    except (AttributeError, IndexError, TypeError, ValueError):
        pass
    return 8640


def python_set_equation_table_layout(table: Any, width_twips: int) -> None:
    """Use a fixed, symmetric 15/70/15 equation table layout."""
    tbl = table._tbl
    tblpr = tbl.tblPr
    tbl_width = python_ensure_child(tblpr, python_qn("w:tblW"), 0)
    tbl_width.set(python_qn("w:type"), "dxa")
    tbl_width.set(python_qn("w:w"), str(width_twips))
    layout = python_ensure_child(tblpr, python_qn("w:tblLayout"))
    layout.set(python_qn("w:type"), "fixed")

    grid = python_child(tbl, python_qn("w:tblGrid"))
    if grid is None:
        grid = python_element("w:tblGrid")
        tbl.insert(tbl.index(tblpr) + 1, grid)

    left_width = int(round(width_twips * 0.15))
    center_width = int(round(width_twips * 0.70))
    column_widths = (left_width, center_width, width_twips - left_width - center_width)

    grid_columns = list(grid.findall(python_qn("w:gridCol")))
    for extra in grid_columns[len(column_widths):]:
        grid.remove(extra)
    while len(grid_columns) < len(column_widths):
        column = python_element("w:gridCol")
        grid.append(column)
        grid_columns.append(column)
    for column, width in zip(grid_columns, column_widths):
        column.set(python_qn("w:w"), str(width))

    for cell, width in zip(table.rows[0].cells, column_widths):
        tcpr = cell._tc.get_or_add_tcPr()
        tcw = python_ensure_child(tcpr, python_qn("w:tcW"), 0)
        tcw.set(python_qn("w:type"), "dxa")
        tcw.set(python_qn("w:w"), str(width))


def python_make_equation_layout(
    doc: Any,
    paragraph: Any,
    math_para: Any,
    number: str,
    style: Dict[str, Any],
) -> bool:
    body = doc._body._body
    if paragraph._p.getparent() is not body:
        return False

    _, WD_TABLE_ALIGNMENT, WD_ALIGN_PARAGRAPH, _, _, _, _, _ = python_docx_api()
    index = body.index(paragraph._p)
    table = doc.add_table(rows=1, cols=3)
    body.remove(table._tbl)
    body.insert(index, table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    python_remove_table_borders(table)
    python_set_equation_table_layout(table, python_equation_table_width_twips(doc))

    center = table.cell(0, 1).paragraphs[0]
    center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    python_apply_paragraph_style(center, style)
    center._p.append(deepcopy(math_para))

    right = table.cell(0, 2).paragraphs[0]
    python_apply_paragraph_style(right, style, preserve_indent=True)
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    number_run = right.add_run(number)
    python_apply_run_style(number_run, style)

    body.remove(paragraph._p)
    return True


def python_format_equation(doc: Any, paragraph: Any, style: Dict[str, Any]) -> None:
    python_apply_paragraph_style(paragraph, style)
    python_apply_runs(paragraph, style)
    python_apply_math_style(paragraph, style)
    math_para = paragraph._p.find(".//" + python_qn("m:oMathPara"))
    if math_para is None:
        return
    number = python_extract_equation_number(math_para)
    if number:
        python_make_equation_layout(doc, paragraph, math_para, number, style)


def format_docx_python(
    input_path: str,
    output_path: str,
    raw_config: Dict[str, Any],
    scope: str = "all",
) -> Dict[str, Any]:
    Document, _, _, _, _, _, Cm, _ = python_docx_api()
    scope = normalize_scope(scope)
    config = deep_merge(default_config(), raw_config or {})
    summary = inspect_docx(input_path)
    doc = Document(input_path)
    if scope in ("all", "equation"):
        python_set_document_math_font(doc, config.get("equation", {}))

    if scope in ("all", "page"):
        for section in doc.sections:
            page = config.get("page", {})
            paper = str(page.get("paper_size", "A4")).upper()
            if paper == "LETTER":
                section.page_width = Cm(21.59)
                section.page_height = Cm(27.94)
            else:
                section.page_width = Cm(21.0)
                section.page_height = Cm(29.7)
            section.top_margin = Cm(to_float(page.get("margin_top", 2.54), 2.54))
            section.bottom_margin = Cm(to_float(page.get("margin_bottom", 2.54), 2.54))
            section.left_margin = Cm(to_float(page.get("margin_left", 3.18), 3.18))
            section.right_margin = Cm(to_float(page.get("margin_right", 3.18), 3.18))

    original_tables = list(doc.tables)
    in_reference_section = False
    reference_number = 0
    formatted_count = 0
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        kind, level = python_classify_paragraph(paragraph)
        reference_heading = is_reference_heading(text)
        if reference_heading and level is None:
            kind, level = "heading", 1
        if in_reference_section and text and not reference_heading and level is not None:
            in_reference_section = False
        if in_reference_section and text and not reference_heading:
            kind, level = "reference", None
        selected = scope_matches(scope, kind, level)
        if selected:
            if kind == "equation":
                python_format_equation(doc, paragraph, config.get("equation", {}))
            elif kind == "heading":
                style = config.get(f"heading{level}", config.get("heading3", {}))
                python_apply_paragraph_style(paragraph, style)
                python_apply_runs(paragraph, style)
            elif kind == "code":
                style = config.get("code", {})
                python_apply_paragraph_style(paragraph, style)
                python_apply_runs(paragraph, style)
            elif kind == "list":
                style = config.get("list", config.get("body", {}))
                python_apply_paragraph_style(paragraph, style, preserve_indent=True)
                python_apply_runs(paragraph, style)
            elif kind == "caption":
                style = config.get("caption", {})
                python_apply_paragraph_style(paragraph, style)
                python_apply_runs(paragraph, style)
            elif kind == "note":
                style = config.get("note", config.get("caption", {}))
                python_apply_paragraph_style(paragraph, style)
                python_apply_runs(paragraph, style)
            elif kind == "reference":
                reference_number += 1
                style = config.get("reference", {})
                python_apply_paragraph_style(paragraph, style)
                python_prepend_reference_number(paragraph, reference_number)
                python_apply_runs(paragraph, style)
            elif kind == "image":
                python_apply_paragraph_style(paragraph, config.get("image", {}), preserve_indent=True)
            else:
                style = config.get("body", {})
                python_apply_paragraph_style(paragraph, style)
                python_apply_runs(paragraph, style)
            formatted_count += 1

        if kind != "equation" and python_has_math(paragraph) and scope in ("all", "equation"):
            python_apply_math_style(paragraph, config.get("equation", {}))
            if scope == "equation" and not selected:
                formatted_count += 1
        if reference_heading:
            in_reference_section = True
            reference_number = 0

    if scope in ("all", "table"):
        table_style_id = python_style_template(doc, str(config.get("table", {}).get("style_id", "af2")))
        for table in original_tables:
            python_format_table(table, config.get("table", {}), table_style_id)
            formatted_count += 1

    atomic_write_docx(output_path, doc.save)
    counts = dict(summary.get("counts", {}))
    counts["formatted"] = formatted_count
    return {
        "counts": counts,
        "warnings": [],
        "output_path": os.path.abspath(output_path),
        "scope": scope,
    }


def response_for(request: Dict[str, Any]) -> Dict[str, Any]:
    action = request.get("action")
    engine = str(request.get("engine") or "python-docx").lower()
    if action == "inspect":
        summary = inspect_docx(str(request.get("input_path", "")))
        return {"ok": True, "action": action, "engine": engine, "summary": summary}
    if action == "format":
        input_path = str(request.get("input_path", ""))
        output_path = str(request.get("output_path", ""))
        raw_config = request.get("config") or {}
        scope = normalize_scope(request.get("scope"))
        if engine in ("native", "native-ooxml", "self", "self-developed"):
            result = format_docx_native(input_path, output_path, raw_config, scope=scope)
        elif engine in ("python-docx", "python_docx", "docx"):
            result = format_docx_python(input_path, output_path, raw_config, scope=scope)
        else:
            raise ValueError(f"不支持的排版引擎：{engine}")
        return {
            "ok": True,
            "action": action,
            "engine": engine,
            "summary": {"counts": result.get("counts", {}), "samples": []},
            "warnings": result.get("warnings", []),
            "output_path": result.get("output_path", ""),
            "scope": result.get("scope", scope),
        }
    if action == "defaults":
        return {"ok": True, "action": action, "config": default_config()}
    raise ValueError(f"不支持的操作：{action}")


def main() -> int:
    line = sys.stdin.readline()
    if not line:
        return 1
    try:
        request = json.loads(line)
        result = response_for(request)
        print(json.dumps(result, ensure_ascii=False), flush=True)
        return 0
    except Exception as exc:  # The client turns this into an in-app error state.
        print(
            json.dumps(
                {"ok": False, "error": str(exc), "error_type": type(exc).__name__},
                ensure_ascii=False,
            ),
            flush=True,
        )
        return 0


if __name__ == "__main__":
    raise SystemExit(main())
